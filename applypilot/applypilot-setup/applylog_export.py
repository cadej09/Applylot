#!/usr/bin/env python3
"""
ApplyPilot -> Master Job Tracker / Application Log exporter.

Reads ApplyPilot's SQLite database and writes a ranked spreadsheet (Excel + CSV)
of your jobs, with fit score, company, an auto-apply-vs-by-hand flag, whether a
tailored resume + cover letter are ready, and a direct apply link.

Common uses:
    python3 applylog_export.py --tracker         # ranked list of all good matches (score >= 7)
    python3 applylog_export.py --tracker --min-score 6
    python3 applylog_export.py --applied         # only jobs you've submitted
    python3 applylog_export.py --all             # every job in the DB
    python3 applylog_export.py --tracker --out ~/Desktop/jobs.xlsx

The database is the source of truth, so this regenerates the whole sheet each run
(no duplicates). Re-run it anytime to refresh.
"""

import argparse
import csv
import os
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

DB_PATH = Path.home() / ".applypilot" / "applypilot.db"
DEFAULT_OUT = Path(__file__).resolve().parent.parent / "job_tracker.xlsx"

# (header, db_column_or_None). None columns are derived in build_records().
COLUMNS = [
    ("Fit Score", "fit_score"),
    ("Apply Method", None),
    ("Company", None),
    ("Title", "title"),
    ("Location", "location"),
    ("Salary", "salary"),
    ("Status", "apply_status"),
    ("Resume Tailored", None),
    ("Cover Letter", None),
    ("Applied On", "applied_at"),
    ("Source", "site"),
    ("Date Found", "discovered_at"),
    ("Apply / Job Link", None),
    ("Job Description", None),
    ("Match Keywords / Notes", "score_reasoning"),
]

AGGREGATORS = ("indeed.com", "linkedin.com", "glassdoor.com",
               "ziprecruiter.com", "google.com")
PATH_ATS = ("greenhouse.io", "lever.co", "ashbyhq.com",
            "smartrecruiters.com", "workable.com", "jobvite.com")
SUBDOMAIN_ATS = ("myworkdayjobs.com", "icims.com", "breezy.hr",
                 "bamboohr.com", "applytojob.com")
GENERIC_SKIP = {"www", "jobs", "apply", "careers", "career", "job",
                "boards", "job-boards", "secure", "recruiting", "us"}


def _prettify(slug: str) -> str:
    slug = slug.replace("-", " ").replace("_", " ").strip()
    return " ".join(w if (w.isupper() and len(w) <= 4) else w.capitalize()
                    for w in slug.split())


def derive_company(url: str, application_url: str) -> str:
    """Best-effort company name from a job/application URL."""
    for u in (application_url, url):
        if not u:
            continue
        try:
            parsed = urlparse(u)
        except Exception:
            continue
        host = (parsed.netloc or "").lower()
        segments = [s for s in (parsed.path or "").split("/") if s]
        if not host:
            continue
        for dom in SUBDOMAIN_ATS:
            if host.endswith(dom):
                sub = host.split(".")[0]
                if sub not in GENERIC_SKIP:
                    return _prettify(sub)
        for dom in PATH_ATS:
            if dom in host and segments:
                return _prettify(segments[0])
        if any(agg in host for agg in AGGREGATORS):
            continue
        parts = host.split(".")
        if len(parts) >= 2:
            label = parts[-2]
            if label not in GENERIC_SKIP and len(label) > 1:
                return _prettify(label)
    return ""


def apply_method(row: dict) -> str:
    """Auto-appliable if there's a real ATS application URL; else by hand."""
    app = row.get("application_url")
    site = (row.get("site") or "").lower()
    if app and not any(agg in app.lower() for agg in AGGREGATORS):
        return "Auto (ATS)"
    if site in ("linkedin", "indeed", "glassdoor"):
        return "By hand"
    return "By hand"


def _fmt_date(val: str) -> str:
    if not val:
        return ""
    try:
        return datetime.fromisoformat(val.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M")
    except Exception:
        return str(val)


def fetch_rows(conn, mode, min_score):
    conn.row_factory = sqlite3.Row
    where = []
    if mode == "applied":
        where.append("applied_at IS NOT NULL")
    elif mode == "tracker":
        where.append(f"COALESCE(fit_score, 0) >= {int(min_score)}")
    # mode == "all": no filter
    sql = "SELECT * FROM jobs"
    if where:
        sql += " WHERE " + " AND ".join(where)
    if mode == "applied":
        sql += " ORDER BY applied_at DESC"
    else:
        sql += " ORDER BY COALESCE(fit_score, 0) DESC, discovered_at DESC"
    return conn.execute(sql).fetchall()


def build_records(rows):
    records = []
    for r in rows:
        r = dict(r)
        company = derive_company(r.get("url", ""), r.get("application_url", ""))
        desc = r.get("full_description") or r.get("description") or ""
        link = r.get("application_url") or r.get("url") or ""
        rec = {}
        for header, col in COLUMNS:
            if header == "Company":
                rec[header] = company
            elif header == "Apply Method":
                rec[header] = apply_method(r)
            elif header == "Resume Tailored":
                # Show the actual résumé file used, so an applied job is traceable
                trp = r.get("tailored_resume_path") or ""
                rec[header] = os.path.basename(trp).rsplit(".", 1)[0] if trp else ""
            elif header == "Cover Letter":
                clp = r.get("cover_letter_path") or ""
                rec[header] = os.path.basename(clp).rsplit(".", 1)[0] if clp else ""
            elif header == "Apply / Job Link":
                rec[header] = link
            elif header == "Job Description":
                rec[header] = desc
            elif col in ("applied_at", "discovered_at"):
                rec[header] = _fmt_date(r.get(col, ""))
            else:
                rec[header] = r.get(col, "")
        records.append(rec)
    return records


def write_csv(records, path):
    headers = [h for h, _ in COLUMNS]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        for rec in records:
            w.writerow(rec)


def write_xlsx(records, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False

    headers = [h for h, _ in COLUMNS]
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for c, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)

    auto_fill = PatternFill("solid", fgColor="E2EFDA")   # green-ish
    hand_fill = PatternFill("solid", fgColor="FCE4D6")   # orange-ish
    method_col = headers.index("Apply Method") + 1

    for row_idx, rec in enumerate(records, start=2):
        for c, header in enumerate(headers, 1):
            val = rec.get(header, "")
            if header == "Job Description" and isinstance(val, str) and len(val) > 32000:
                val = val[:32000] + " ...[truncated]"
            cell = ws.cell(row=row_idx, column=c, value=val)
        # Color the Apply Method cell
        mcell = ws.cell(row=row_idx, column=method_col)
        mcell.fill = auto_fill if str(mcell.value).startswith("Auto") else hand_fill

    widths = {
        "Fit Score": 8, "Apply Method": 12, "Company": 22, "Title": 32,
        "Location": 20, "Salary": 16, "Status": 11, "Resume Tailored": 10,
        "Cover Letter": 10, "Applied On": 16, "Source": 10, "Date Found": 16,
        "Apply / Job Link": 42, "Job Description": 60, "Match Keywords / Notes": 40,
    }
    for c, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(header, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    wb.save(path)
    return True


def main():
    ap = argparse.ArgumentParser(description="Export ApplyPilot jobs to a ranked Excel/CSV tracker.")
    ap.add_argument("--db", default=str(DB_PATH), help="Path to applypilot.db")
    ap.add_argument("--out", default=str(DEFAULT_OUT), help="Output .xlsx path")
    group = ap.add_mutually_exclusive_group()
    group.add_argument("--tracker", action="store_true",
                       help="Ranked list of all matches scoring >= --min-score (default mode)")
    group.add_argument("--applied", action="store_true", help="Only jobs you've applied to")
    group.add_argument("--all", action="store_true", help="Every job in the DB")
    ap.add_argument("--min-score", type=int, default=7, help="Score threshold for --tracker (default 7)")
    args = ap.parse_args()

    mode = "applied" if args.applied else ("all" if args.all else "tracker")

    db_path = Path(args.db).expanduser()
    if not db_path.exists():
        sys.exit(f"Database not found at {db_path}. Has ApplyPilot run yet?")

    out_xlsx = Path(args.out).expanduser()
    out_csv = out_xlsx.with_suffix(".csv")

    conn = sqlite3.connect(str(db_path))
    rows = fetch_rows(conn, mode, args.min_score)
    records = build_records(rows)
    conn.close()

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    write_csv(records, out_csv)
    xlsx_ok = write_xlsx(records, out_xlsx)

    label = {"applied": "applied jobs", "all": "total jobs",
             "tracker": f"matches (score >= {args.min_score})"}[mode]
    auto = sum(1 for r in records if str(r.get("Apply Method", "")).startswith("Auto"))
    print(f"Exported {len(records)} {label}  —  {auto} auto-appliable, {len(records)-auto} by hand.")
    print(f"  CSV : {out_csv}")
    if xlsx_ok:
        print(f"  XLSX: {out_xlsx}")
    else:
        print("  XLSX skipped (openpyxl not installed). Run: python3 -m pip install openpyxl")
    if not records:
        print("\nNote: nothing matched. Try a lower --min-score, or run the pipeline first.")


if __name__ == "__main__":
    main()
